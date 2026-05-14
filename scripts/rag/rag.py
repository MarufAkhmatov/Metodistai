#!/usr/bin/env python3
"""
AI Metodist — bilim bazasi (RAG) skripti.

Buyruqlar:
  python rag.py index "<agent_dir>"     — papkani skanlab indeks va vektorlarni yangilaydi
  python rag.py query "<agent_dir>"     — savol (stdin orqali) bo'yicha tegishli parchalarni JSON qaytaradi
  python rag.py status "<agent_dir>"    — indeks holatini JSON qaytaradi
  python rag.py models                  — fastembed qo'llab-quvvatlaydigan modellar ro'yxati

Natija JSON ko'rinishida stdout'ning oxirgi qatoriga chiqadi. Loglar stderr'ga.
"""
import sys
import os
import json
import time
import traceback
from pathlib import Path
from collections import defaultdict

EMBED_MODEL = os.environ.get(
    "RAG_EMBED_MODEL", "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
)
DOC_EXTS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".md"}
INDEX_MD_NAME = "KNOWLEDGE_INDEX.md"
RAG_DIR_NAME = ".rag"
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
TOP_K = 8

_model = None


def log(*args):
    print("[rag]", *args, file=sys.stderr, flush=True)


def print_result(obj):
    print(json.dumps(obj, ensure_ascii=False))
    sys.stdout.flush()


def load_model():
    global _model
    if _model is None:
        try:
            from fastembed import TextEmbedding
        except ImportError as e:
            raise RuntimeError(
                "fastembed o'rnatilmagan. Ishga tushiring: "
                "pip install -r scripts/rag/requirements.txt"
            ) from e
        log(f"embedding modeli yuklanmoqda: {EMBED_MODEL} (birinchi marta — yuklab olinadi)")
        _model = TextEmbedding(model_name=EMBED_MODEL)
    return _model


# --- Matn ajratish ---

def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            pass
    return "\n".join(parts)


def extract_docx(path):
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def extract_xls(path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    parts = []
    for sheet in book.sheets():
        parts.append(f"# {sheet.name}")
        for r in range(sheet.nrows):
            cells = [str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            cells = [c for c in cells if c.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def extract_text_file(path):
    return path.read_text(encoding="utf-8", errors="ignore")


def extract(path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".xlsx":
        return extract_xlsx(path)
    if ext == ".xls":
        return extract_xls(path)
    if ext in (".txt", ".md"):
        return extract_text_file(path)
    if ext == ".doc":
        raise RuntimeError("eski .doc formati qo'llab-quvvatlanmaydi — .docx ga o'tkazing")
    raise RuntimeError(f"noma'lum fayl turi: {ext}")


def chunk_text(text):
    text = " ".join(text.split())
    if not text:
        return []
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + CHUNK_SIZE, n)
        chunks.append(text[start:end])
        if end >= n:
            break
        start = end - CHUNK_OVERLAP
    return chunks


def iter_doc_files(agent_dir):
    agent_dir = Path(agent_dir)
    for root, dirs, files in os.walk(agent_dir):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for fn in files:
            if fn.startswith(".") or fn == INDEX_MD_NAME:
                continue
            p = Path(root) / fn
            if p.suffix.lower() in DOC_EXTS:
                yield p


def store_paths(agent_dir):
    rag_dir = Path(agent_dir) / RAG_DIR_NAME
    return {
        "dir": rag_dir,
        "chunks": rag_dir / "chunks.json",
        "embeddings": rag_dir / "embeddings.npy",
        "manifest": rag_dir / "manifest.json",
    }


def write_index_md(agent_dir, files, chunks, skipped):
    agent_dir = Path(agent_dir)
    chunks_by_path = defaultdict(int)
    for c in chunks:
        chunks_by_path[c["rel_path"]] += 1

    by_folder = defaultdict(list)
    for p in files:
        rel = str(p.relative_to(agent_dir)).replace("\\", "/")
        folder = os.path.dirname(rel) or "."
        st = p.stat()
        by_folder[folder].append(
            {
                "name": p.name,
                "ext": p.suffix.lower().lstrip("."),
                "size": st.st_size,
                "chunks": chunks_by_path.get(rel, 0),
            }
        )

    lines = [
        "# AI Metodist — Bilim bazasi indeksi",
        "",
        "> Bu fayl avtomatik yaratiladi. Qo'lda tahrirlamang — papka o'zgarganda qayta yoziladi.",
        "",
        f"- Yangilangan: {time.strftime('%Y-%m-%d %H:%M:%S')}",
        f"- Jami hujjatlar: {len(files)}",
        f"- Jami parchalar (chunks): {len(chunks)}",
        f"- Embedding modeli: {EMBED_MODEL}",
        "",
    ]
    for folder in sorted(by_folder.keys()):
        title = folder if folder != "." else "(ildiz papka)"
        lines.append(f"## {title}")
        lines.append("")
        for d in sorted(by_folder[folder], key=lambda x: x["name"].lower()):
            kb = max(1, round(d["size"] / 1024))
            note = "" if d["chunks"] > 0 else "  _(indekslanmagan)_"
            lines.append(
                f"- **{d['name']}** — {d['ext'].upper()}, {kb} KB, {d['chunks']} parcha{note}"
            )
        lines.append("")
    if skipped:
        lines.append("## Indekslanmagan fayllar")
        lines.append("")
        for s in skipped:
            lines.append(f"- {s['file']} — {s['reason']}")
        lines.append("")

    (agent_dir / INDEX_MD_NAME).write_text("\n".join(lines), encoding="utf-8")


def cmd_index(agent_dir):
    import numpy as np

    agent_dir = Path(agent_dir)
    sp = store_paths(agent_dir)
    sp["dir"].mkdir(parents=True, exist_ok=True)

    old_manifest = {}
    old_chunks = []
    old_emb = None
    if sp["manifest"].exists() and sp["chunks"].exists() and sp["embeddings"].exists():
        try:
            old_manifest = json.loads(sp["manifest"].read_text(encoding="utf-8"))
            old_chunks = json.loads(sp["chunks"].read_text(encoding="utf-8"))
            old_emb = np.load(sp["embeddings"])
        except Exception as e:
            log("eski indeksni o'qib bo'lmadi, qaytadan quriladi:", e)
            old_manifest, old_chunks, old_emb = {}, [], None

    old_by_path = defaultdict(list)
    if old_emb is not None and len(old_chunks) == len(old_emb):
        for i, ch in enumerate(old_chunks):
            old_by_path[ch["rel_path"]].append((ch, old_emb[i]))

    files = sorted(iter_doc_files(agent_dir), key=lambda p: str(p).lower())
    new_manifest = {}
    final_chunks = []
    final_emb_rows = []
    to_embed_texts = []
    to_embed_refs = []
    skipped = []

    for p in files:
        rel = str(p.relative_to(agent_dir)).replace("\\", "/")
        st = p.stat()
        sig = {"mtime": int(st.st_mtime), "size": st.st_size}
        prev = old_manifest.get(rel)
        if (
            prev
            and prev.get("mtime") == sig["mtime"]
            and prev.get("size") == sig["size"]
            and rel in old_by_path
        ):
            new_manifest[rel] = sig
            for ch, emb in old_by_path[rel]:
                final_chunks.append(ch)
                final_emb_rows.append(emb)
            continue

        try:
            text = extract(p)
        except Exception as e:
            log(f"o'qib bo'lmadi {rel}: {e}")
            skipped.append({"file": rel, "reason": str(e)})
            continue
        chunks = chunk_text(text)
        if not chunks:
            skipped.append({"file": rel, "reason": "bo'sh matn"})
            continue
        new_manifest[rel] = sig
        for idx, ctext in enumerate(chunks):
            final_chunks.append(
                {"rel_path": rel, "chunk_index": idx, "text": ctext, "mtime": sig["mtime"]}
            )
            final_emb_rows.append(None)
            to_embed_texts.append(ctext)
            to_embed_refs.append(len(final_chunks) - 1)

    if to_embed_texts:
        model = load_model()
        log(f"{len(to_embed_texts)} ta yangi parcha embedding qilinmoqda...")
        vectors = list(model.embed(to_embed_texts))
        for ref, vec in zip(to_embed_refs, vectors):
            final_emb_rows[ref] = np.asarray(vec, dtype=np.float32)

    if final_emb_rows:
        emb_array = np.vstack([np.asarray(r, dtype=np.float32) for r in final_emb_rows])
    else:
        emb_array = np.zeros((0, 384), dtype=np.float32)

    sp["chunks"].write_text(json.dumps(final_chunks, ensure_ascii=False), encoding="utf-8")
    np.save(sp["embeddings"], emb_array)
    sp["manifest"].write_text(
        json.dumps(new_manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    write_index_md(agent_dir, files, final_chunks, skipped)

    print_result(
        {
            "ok": True,
            "indexed_files": len(new_manifest),
            "total_chunks": len(final_chunks),
            "skipped": skipped,
            "model": EMBED_MODEL,
            "updated_at": int(time.time()),
        }
    )


def cmd_query(agent_dir):
    import numpy as np

    query = sys.stdin.read().strip()
    sp = store_paths(agent_dir)
    if not (sp["chunks"].exists() and sp["embeddings"].exists()):
        print_result({"ok": True, "chunks": [], "indexed_files": 0, "note": "indeks hali tayyor emas"})
        return
    if not query:
        print_result({"ok": True, "chunks": [], "indexed_files": 0, "note": "bo'sh so'rov"})
        return

    chunks = json.loads(sp["chunks"].read_text(encoding="utf-8"))
    emb = np.load(sp["embeddings"])
    if len(chunks) == 0 or emb.shape[0] == 0:
        print_result({"ok": True, "chunks": [], "indexed_files": 0})
        return

    model = load_model()
    qvec = np.asarray(list(model.embed([query]))[0], dtype=np.float32)
    emb_norm = emb / (np.linalg.norm(emb, axis=1, keepdims=True) + 1e-9)
    qn = qvec / (np.linalg.norm(qvec) + 1e-9)
    scores = emb_norm @ qn
    top_idx = np.argsort(-scores)[:TOP_K]

    results = []
    for i in top_idx:
        ch = chunks[int(i)]
        results.append(
            {
                "rel_path": ch["rel_path"],
                "chunk_index": ch["chunk_index"],
                "text": ch["text"],
                "score": round(float(scores[int(i)]), 4),
            }
        )

    indexed_files = len({c["rel_path"] for c in chunks})
    print_result({"ok": True, "chunks": results, "indexed_files": indexed_files})


def cmd_status(agent_dir):
    sp = store_paths(agent_dir)
    if not sp["manifest"].exists():
        print_result({"ok": True, "ready": False, "indexed_files": 0, "total_chunks": 0})
        return
    manifest = json.loads(sp["manifest"].read_text(encoding="utf-8"))
    total_chunks = 0
    if sp["chunks"].exists():
        try:
            total_chunks = len(json.loads(sp["chunks"].read_text(encoding="utf-8")))
        except Exception:
            pass
    print_result(
        {
            "ok": True,
            "ready": True,
            "indexed_files": len(manifest),
            "total_chunks": total_chunks,
            "model": EMBED_MODEL,
        }
    )


def main():
    args = sys.argv[1:]
    if not args:
        print_result({"ok": False, "error": "buyruq ko'rsatilmadi"})
        sys.exit(1)
    cmd = args[0]
    try:
        if cmd == "index":
            cmd_index(args[1])
        elif cmd == "query":
            cmd_query(args[1])
        elif cmd == "status":
            cmd_status(args[1])
        elif cmd == "models":
            from fastembed import TextEmbedding

            models = [m.get("model") for m in TextEmbedding.list_supported_models()]
            print_result({"ok": True, "models": models})
        else:
            print_result({"ok": False, "error": f"noma'lum buyruq: {cmd}"})
            sys.exit(1)
    except IndexError:
        print_result({"ok": False, "error": "papka yo'li ko'rsatilmadi"})
        sys.exit(1)
    except Exception as e:
        log("XATO:", traceback.format_exc())
        print_result({"ok": False, "error": str(e)})
        sys.exit(1)


if __name__ == "__main__":
    main()
